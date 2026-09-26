# SPDX-License-Identifier: MIT
"""Tests der Extraktionsschicht — Schwerpunkt: der Fix fuer verbundene Zellen.

Der erste Test ist die Regression zu Ticket T-20260817-05. Er baut eine DOCX
mit einer ueber mehrere Spalten verbundenen Zelle und prueft, dass der Inhalt
GENAU EINMAL erscheint. Ohne den Fix erschiene er einmal je Rasterspalte
(an einer realen Datei gemessen: Faktor 17,8).
"""

from __future__ import annotations

import pytest

from doc_services.extract import (
    SUFFIX_MAP,
    ExtraktionsFehler,
    _zellen_einer_zeile,
    extrahieren,
    format_von,
)

docx = pytest.importorskip("docx", reason="python-docx wird fuer diese Tests gebraucht")


def _docx_mit_verbundener_zelle(ziel, spalten: int = 20):
    """Erzeugt eine DOCX, deren erste Zeile ueber alle Spalten verbunden ist."""
    from docx import Document
    d = Document()
    t = d.add_table(rows=2, cols=spalten)
    verbunden = t.cell(0, 0).merge(t.cell(0, spalten - 1))
    verbunden.text = "LERNVEREINBARUNG"
    t.cell(1, 0).text = "Vertragspartner"
    t.cell(1, 1).text = "Beginn"
    d.save(str(ziel))
    return ziel


def test_verbundene_zelle_erscheint_nur_einmal(tmp_path):
    """Regression T-20260817-05: keine Vervielfachung verbundener Zellen."""
    from docx import Document

    pfad = _docx_mit_verbundener_zelle(tmp_path / "verbund.docx", spalten=20)
    d = Document(str(pfad))
    zeile = d.tables[0].rows[0]

    # Ausgangslage bestaetigen: python-docx liefert die Zelle mehrfach
    roh = [c.text.strip() for c in zeile.cells if c.text.strip()]
    assert len(roh) > 1, "Testaufbau taugt nicht - Zelle wurde nicht vervielfacht"

    # Der Fix fasst sie zusammen
    entdoppelt = _zellen_einer_zeile(zeile)
    assert entdoppelt == ["LERNVEREINBARUNG"], (
        f"Verbundene Zelle wurde nicht dedupliziert: {entdoppelt}"
    )


def test_vervielfachungsfaktor_wird_aufgehoben(tmp_path):
    """Der gemessene Faktor darf nicht wieder auftauchen."""
    from docx import Document

    spalten = 20
    pfad = _docx_mit_verbundener_zelle(tmp_path / "faktor.docx", spalten=spalten)
    zeile = Document(str(pfad)).tables[0].rows[0]

    naiv = len([c.text.strip() for c in zeile.cells if c.text.strip()])
    fix = len(_zellen_einer_zeile(zeile))
    assert fix == 1
    assert naiv / fix >= 2, "Testaufbau erzeugt keine messbare Vervielfachung"


def test_nicht_verbundene_zellen_bleiben_vollstaendig(tmp_path):
    """Der Fix darf keine echten Inhalte schlucken."""
    from docx import Document

    pfad = _docx_mit_verbundener_zelle(tmp_path / "normal.docx", spalten=20)
    zeile = Document(str(pfad)).tables[0].rows[1]
    assert _zellen_einer_zeile(zeile) == ["Vertragspartner", "Beginn"]


def test_format_erkennung():
    assert format_von("x.docx") == "docx"
    assert format_von("X.PDF") == "pdf"
    assert format_von("a.md") == "txt"
    assert format_von("b.unbekannt") is None


def test_alle_suffixe_haben_katalogeintrag():
    """Jedes gemappte Format muss im Katalog stehen, sonst laeuft es ins Leere."""
    from doc_services.config import Registry

    reg = Registry()
    bekannt = set(reg.formate())
    fehlend = {v for v in SUFFIX_MAP.values()} - bekannt
    assert not fehlend, f"Formate ohne Katalogeintrag: {sorted(fehlend)}"


def test_unbekannte_endung_wird_klar_gemeldet(tmp_path):
    p = tmp_path / "datei.xyz"
    p.write_text("inhalt", encoding="utf-8")
    with pytest.raises(ExtraktionsFehler, match="Unbekannte Endung"):
        extrahieren(p)


def test_fehlende_datei(tmp_path):
    with pytest.raises(FileNotFoundError):
        extrahieren(tmp_path / "gibtsnicht.docx")


def test_textdatei_wird_direkt_gelesen(tmp_path):
    p = tmp_path / "notiz.md"
    p.write_text("# Überschrift\n\nMit echten Umlauten: äöüß\n", encoding="utf-8")
    e = extrahieren(p, lernen=False)
    assert "Überschrift" in e.text
    assert "äöüß" in e.text, "Umlaute muessen unversehrt durchkommen"
    assert e.backend == "direkt"
