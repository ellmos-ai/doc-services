# SPDX-License-Identifier: MIT
"""Dokument -> Text/Markdown, mit Praeferenzkette und ehrlichem Fallback.

Herkunft: portiert aus BACH `system/hub/_services/document/document_pipeline.py`
(MIT, eigenes Copyright) und dort BACH-unabhaengig gemacht - keine Importe aus
`hub`, `tools` oder `bach_paths`, keine SQLite-Kopplung, keine Klientenlogik.

Zwei Dinge sind gegenueber der BACH-Fassung bewusst anders:

1. BEHOBENER FEHLER - verbundene Tabellenzellen.
   python-docx liefert eine ueber mehrere Spalten verbundene Zelle EINMAL PRO
   RASTERSPALTE. Die BACH-Fassung dedupliziert nicht; gemessen am 2026-08-17 an
   einer realen 25-KB-Datei: 785 gezaehlte gegen 44 tatsaechlich eindeutige
   Zellen (Faktor 17,8), 44.670 unnoetige Zeichen, rund 12.760 Tokens. Hier wird
   ueber die Identitaet des zugrunde liegenden tc-Elements dedupliziert.
   Belegt in Ticket T-20260817-825816579.

2. KEINE AGPL-BINDUNG.
   PDF-Rendering laeuft ausschliesslich ueber pypdfium2 (BSD-3/Apache-2.0).
   PyMuPDF/fitz wird nirgends importiert (Entscheidung E08 vom 2026-08-18).
"""

from __future__ import annotations

import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

from .config import Registry
from .config import registry as _registry

# Ausgabeformen, die ein Backend liefern kann (Feld `produces` im Katalog).
AUSGABEFORMEN: frozenset[str] = frozenset({"text", "markdown"})

# Endung -> Formatschluessel im Katalog
SUFFIX_MAP: dict[str, str] = {
    ".docx": "docx", ".doc": "doc",
    ".xlsx": "xlsx", ".xls": "xls",
    ".pdf": "pdf", ".pptx": "pptx",
    ".msg": "msg", ".eml": "eml",
    ".epub": "epub", ".html": "html", ".htm": "html",
    ".csv": "csv",
    ".txt": "txt", ".md": "txt", ".markdown": "txt",
    ".jpg": "image", ".jpeg": "image", ".png": "image",
    ".tif": "image", ".tiff": "image", ".bmp": "image",
}


class ExtraktionsFehler(Exception):
    """Kein Backend konnte die Datei lesen."""


@dataclass
class Ergebnis:
    """Was bei der Extraktion herauskam - samt Weg dorthin."""
    pfad: Path
    format: str
    text: str
    backend: str
    produces: str                       # "markdown" | "text"
    versuche: list[tuple[str, str]] = field(default_factory=list)  # (backend, ergebnis)
    hinweise: list[str] = field(default_factory=list)

    @property
    def ist_markdown(self) -> bool:
        return self.produces == "markdown"

    @property
    def zeichen(self) -> int:
        return len(self.text)


# ----------------------------------------------------------------------------
# Einzelne Backends
# ----------------------------------------------------------------------------

def _b_markitdown(pfad: Path) -> str:
    from markitdown import MarkItDown
    ergebnis = MarkItDown().convert(str(pfad))
    return ergebnis.text_content or ""


def _zellen_einer_zeile(row) -> list[str]:
    """Zellen einer Tabellenzeile OHNE Vervielfachung verbundener Zellen.

    Der Kern des Fixes: `row.cells` gibt dieselbe verbundene Zelle einmal je
    Rasterspalte zurueck. Ueber die Identitaet des tc-XML-Elements laesst sich
    das zuverlaessig zusammenfassen - `id(cell)` genuegt NICHT, weil python-docx
    bei jedem Zugriff neue Wrapper-Objekte erzeugt.
    """
    gesehen: set[int] = set()
    raus: list[str] = []
    for zelle in row.cells:
        schluessel = id(zelle._tc)
        if schluessel in gesehen:
            continue
        gesehen.add(schluessel)
        txt = zelle.text.strip()
        if txt:
            raus.append(txt)
    return raus


def _b_python_docx(pfad: Path) -> str:
    from docx import Document
    doc = Document(str(pfad))
    teile: list[str] = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for tabelle in doc.tables:
        for zeile in tabelle.rows:
            zellen = _zellen_einer_zeile(zeile)
            if zellen:
                teile.append(" | ".join(zellen))
    return "\n".join(teile)


def _b_openpyxl(pfad: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(str(pfad), data_only=True, read_only=True)
    teile: list[str] = []
    for ws in wb.worksheets:
        teile.append(f"## {ws.title}")
        for zeile in ws.iter_rows(values_only=True):
            werte = [str(z).strip() for z in zeile if z is not None and str(z).strip()]
            if werte:
                teile.append(" | ".join(werte))
    wb.close()
    return "\n".join(teile)


def _b_pypdf(pfad: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(pfad))
    return "\n".join(seite.extract_text() or "" for seite in reader.pages)


def _b_ocr(pfad: Path) -> str:
    from .ocr import ocr_datei
    return ocr_datei(pfad)


def _b_antiword(pfad: Path) -> str:
    r = subprocess.run(["antiword", str(pfad)], capture_output=True, text=True,
                       encoding="utf-8", errors="replace", timeout=60)
    if r.returncode != 0 or not r.stdout.strip():
        raise ExtraktionsFehler(f"antiword lieferte nichts (rc={r.returncode})")
    return r.stdout.strip()


def _b_libreoffice(pfad: Path) -> str:
    with tempfile.TemporaryDirectory() as tmp:
        r = subprocess.run(
            ["soffice", "--headless", "--convert-to", "txt:Text", "--outdir", tmp, str(pfad)],
            capture_output=True, timeout=180,
        )
        ziel = Path(tmp) / (pfad.stem + ".txt")
        if r.returncode != 0 or not ziel.exists():
            raise ExtraktionsFehler(f"LibreOffice-Konvertierung fehlgeschlagen (rc={r.returncode})")
        return ziel.read_text(encoding="utf-8", errors="replace")


def _b_direkt(pfad: Path) -> str:
    for kodierung in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return pfad.read_text(encoding=kodierung)
        except UnicodeDecodeError:
            continue
    return pfad.read_text(encoding="utf-8", errors="replace")


def _b_stdlib_email(pfad: Path) -> str:
    import email
    from email import policy
    nachricht = email.message_from_bytes(pfad.read_bytes(), policy=policy.default)
    kopf = [f"{k}: {v}" for k, v in nachricht.items()
            if k.lower() in {"from", "to", "cc", "subject", "date"}]
    teil = nachricht.get_body(preferencelist=("plain", "html"))
    rumpf = teil.get_content() if teil else ""
    return "\n".join(kopf) + "\n\n" + rumpf


def _b_stdlib_csv(pfad: Path) -> str:
    import csv as _csv
    import io
    roh = _b_direkt(pfad)
    try:
        dialekt = _csv.Sniffer().sniff(roh[:4096])
    except _csv.Error:
        dialekt = _csv.excel
    zeilen = ["|".join(z) for z in _csv.reader(io.StringIO(roh), dialekt)]
    return "\n".join(zeilen)


BACKENDS = {
    "markitdown": _b_markitdown,
    "python-docx-native": _b_python_docx,
    "openpyxl-native": _b_openpyxl,
    "pypdf-native": _b_pypdf,
    "ocr-tesseract": _b_ocr,
    "antiword": _b_antiword,
    "libreoffice": _b_libreoffice,
    "direkt": _b_direkt,
    "stdlib-email": _b_stdlib_email,
    "stdlib-csv": _b_stdlib_csv,
}


# ----------------------------------------------------------------------------
# Oeffentliche Schnittstelle
# ----------------------------------------------------------------------------

def format_von(pfad: Path | str) -> str | None:
    return SUFFIX_MAP.get(Path(pfad).suffix.lower())


def extrahieren(
    pfad: Path | str,
    reg: Registry | None = None,
    nur_backend: str | None = None,
    lernen: bool = True,
    produces: str | None = None,
) -> Ergebnis:
    """Extrahiert eine Datei ueber die konfigurierte Praeferenzkette.

    Args:
        pfad: Quelldatei.
        reg: Registry; ohne Angabe wird die Standardkonfiguration geladen.
        nur_backend: Erzwingt ein bestimmtes Backend (fuer Vergleichsmessungen).
        lernen: Ergebnis in den Statistiken mitzaehlen.
        produces: Verlangt eine Ausgabeform -- "text" oder "markdown". Die
            Praeferenzkette bleibt erhalten, sie wird nur auf Backends dieser
            Form eingeengt (inklusive Fallbacks). Ohne Angabe aendert sich
            nichts: es gilt die konfigurierte Kette wie bisher.

            Wofuer das da ist: Wer den Text weiterverarbeitet statt ihn
            anzuzeigen -- Schwaerzung, Suche, Regeln auf Wortgrenzen --, braucht
            Fliesstext und keine Markdown-Auszeichnung. Bei LaTeX-PDFs ist das
            keine Geschmacksfrage: markitdown verliert dort Wortabstaende
            (gemessen 2026-09-13 an einem echten Paper: 8.698 statt 14.856
            Woerter, und 37 % der Tokens existieren in der .tex-Quelle gar
            nicht, weil ganze Saetze zu einem Token verkleben). Mit
            produces="text" liefert dieselbe Funktion an derselben Datei ein
            Ergebnis, das mit pypdf Byte fuer Byte uebereinstimmt.

    Raises:
        ExtraktionsFehler: wenn KEIN Backend die Datei lesen konnte - mit der
            vollstaendigen Liste der Versuche und ihrer Fehler.
    """
    pfad = Path(pfad)
    if not pfad.exists():
        raise FileNotFoundError(f"Datei nicht gefunden: {pfad}")

    fmt = format_von(pfad)
    if fmt is None:
        raise ExtraktionsFehler(
            f"Unbekannte Endung '{pfad.suffix}'. Bekannt: {sorted(set(SUFFIX_MAP))}"
        )

    if produces is not None and produces not in AUSGABEFORMEN:
        raise ExtraktionsFehler(
            f"produces={produces!r} ist unbekannt. Erlaubt: "
            + ", ".join(sorted(AUSGABEFORMEN))
            + ". Ohne Angabe gilt die konfigurierte Kette."
        )

    reg = reg or _registry()
    auswahl = reg.aufloesen(fmt)
    kette = [
        b for b in auswahl.kette
        if (nur_backend is None or b.id == nur_backend)
        and (produces is None or b.produces == produces)
    ]

    if not kette:
        ausgeschlossen = "; ".join(f"{x.id}: {x.begruendung}" for x in auswahl.ausgeschlossen)
        verlangt = []
        if nur_backend is not None:
            verlangt.append(f"nur_backend={nur_backend!r}")
        if produces is not None:
            verlangt.append(f"produces={produces!r}")
        raise ExtraktionsFehler(
            f"Kein verfuegbares Backend fuer '.{fmt}'"
            + (f" mit {' und '.join(verlangt)}" if verlangt else "")
            + ". "
            + (f"Ausgeschlossen wurden: {ausgeschlossen}" if ausgeschlossen else
               "Im Katalog ist fuer dieses Format nichts hinterlegt.")
        )

    versuche: list[tuple[str, str]] = []
    for lage in kette:
        fn = BACKENDS.get(lage.id)
        if fn is None:
            versuche.append((lage.id, "nicht implementiert"))
            continue
        try:
            text = fn(pfad)
        except Exception as exc:
            versuche.append((lage.id, f"{type(exc).__name__}: {exc}"))
            if lernen:
                reg.merke_ergebnis(fmt, lage.id, erfolg=False)
            continue

        if not text.strip():
            versuche.append((lage.id, "leeres Ergebnis"))
            if lernen:
                reg.merke_ergebnis(fmt, lage.id, erfolg=False)
            continue

        versuche.append((lage.id, "ok"))
        if lernen:
            reg.merke_ergebnis(fmt, lage.id, erfolg=True)

        hinweise: list[str] = []
        if fmt == "pdf" and lage.id == "markitdown":
            hinweise.append(
                "LaTeX-PDFs verlieren bei diesem Weg Wortabstaende (belegt 2026-08-17). "
                "Bei eigenen Papers die .tex-Quelle nutzen statt das PDF."
            )
        if len(versuche) > 1:
            hinweise.append(
                "Bevorzugtes Backend hat nicht getragen; Fallback wurde verwendet."
            )
        return Ergebnis(pfad=pfad, format=fmt, text=text, backend=lage.id,
                        produces=lage.produces, versuche=versuche, hinweise=hinweise)

    bericht = "; ".join(f"{b}: {e}" for b, e in versuche)
    raise ExtraktionsFehler(f"Alle Backends fuer '.{fmt}' sind gescheitert - {bericht}")
